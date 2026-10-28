"""
Stage 1 — Caption sources.

WHY THIS EXISTS
    This is the seam between "where our supervision physically lives" and
    "everything else". A CaptionSource answers one question:

        give me every validated (class, image number, caption) record

    and nothing downstream knows whether that came from a spreadsheet, a CSV, a
    review tool, or a second crop's annotation round. Adding a new dataset is a
    new class in this file plus one line in SOURCES — no change anywhere else in
    the pipeline.

INPUT   an annotation artifact on disk (today: Revised_captions.xlsx)
OUTPUT  a flat list of CaptionRecord, plus a hash of the artifact for lineage

WHAT A SOURCE MUST NOT DO
    Write captions. Rewrite captions. Infer captions. A source reads what humans
    validated and reports exactly that, or raises.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Protocol

from ..config import Config
from ..io import sha256_file


@dataclass(frozen=True)
class CaptionRecord:
    """One validated caption, still tied to its spreadsheet cell for traceability."""
    disease: str
    image_no: int
    caption_index: int
    caption: str
    origin: str          # e.g. "C7" — where a human can go to check it


class CaptionSource(Protocol):
    """Contract every annotation source implements."""

    name: str

    def load(self) -> List[CaptionRecord]:
        """Return every validated caption. Raise on any structural problem."""
        ...

    def fingerprint(self) -> Dict[str, str]:
        """Identify the exact artifact that was read, for the lineage record."""
        ...


# ── Caption normalisation ────────────────────────────────────────────────────
# The ONLY transformation ever applied to a validated caption.

_PERIOD_GLUE = re.compile(r"\.([A-Z])")
_WHITESPACE = re.compile(r"\s+")


def normalize_whitespace(text: str) -> str:
    """
    Formatting only. Restores the space after a sentence-ending period where two
    sentences were concatenated during copy-paste, and collapses repeated
    whitespace. No word is added, removed, reordered, or rewritten.
    """
    return _WHITESPACE.sub(" ", _PERIOD_GLUE.sub(r". \1", text)).strip()


# ── Excel ────────────────────────────────────────────────────────────────────

class ExcelCaptionSource:
    """
    Reads a wide-format workbook: one row per image, one column per caption.

    The disease column is only filled on the first row of each block, so the last
    non-empty value is carried forward. Every structural surprise is a hard error
    with the offending cell named — a silent misread here would poison the entire
    downstream pipeline.
    """

    name = "excel"

    def __init__(self, cfg: Config) -> None:
        self.cfg = cfg
        self.src = cfg.data.source
        self.path = cfg.paths.excel

    def fingerprint(self) -> Dict[str, str]:
        return {
            "type": self.name,
            "path": self.path.as_posix(),
            "sha256": sha256_file(self.path),
        }

    def load(self) -> List[CaptionRecord]:
        import openpyxl

        if not self.path.exists():
            raise FileNotFoundError(f"caption workbook not found: {self.path}")

        wb = openpyxl.load_workbook(self.path, data_only=True)
        if self.src.sheet not in wb.sheetnames:
            raise ValueError(
                f"sheet '{self.src.sheet}' not in {wb.sheetnames}")
        ws = wb[self.src.sheet]

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ValueError(f"{self.path} is empty")

        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        expected = ([self.src.disease_column, self.src.image_column]
                    + list(self.src.caption_columns))
        if header[:len(expected)] != expected:
            raise ValueError(
                f"unexpected header in {self.path.name}\n"
                f"  found:    {header[:len(expected)]}\n"
                f"  expected: {expected}")

        caption_start = 2
        records: List[CaptionRecord] = []
        current: str | None = None

        for excel_row, row in enumerate(rows[1:], start=2):
            if all(cell is None for cell in row):
                continue

            if isinstance(row[0], str) and row[0].strip():
                current = row[0].strip()
            if current is None:
                raise ValueError(
                    f"row {excel_row}: captions present before any disease name")
            if current not in self.src.class_map:
                raise ValueError(
                    f"row {excel_row}: class '{current}' is not in class_map. "
                    f"Add it to the config deliberately — classes are never "
                    f"auto-discovered.")

            if row[1] is None:
                raise ValueError(f"row {excel_row}: missing image number")
            image_no = int(float(row[1]))

            for offset in range(len(self.src.caption_columns)):
                cell = row[caption_start + offset]
                if cell is None or not str(cell).strip():
                    raise ValueError(
                        f"row {excel_row}: '{current}' image {image_no} "
                        f"{self.src.caption_columns[offset]} is empty")
                records.append(CaptionRecord(
                    disease=current,
                    image_no=image_no,
                    caption_index=offset + 1,
                    caption=normalize_whitespace(str(cell)),
                    origin=f"{chr(ord('C') + offset)}{excel_row}",
                ))

        return records


# ── Registry ─────────────────────────────────────────────────────────────────

SOURCES: Dict[str, type] = {
    ExcelCaptionSource.name: ExcelCaptionSource,
}


def get_source(cfg: Config) -> CaptionSource:
    kind = cfg.data.source.type
    if kind not in SOURCES:
        raise ValueError(
            f"unknown caption source '{kind}'. Available: {sorted(SOURCES)}")
    return SOURCES[kind](cfg)
